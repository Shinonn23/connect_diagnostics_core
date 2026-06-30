# Copyright (c) 2026, Siwat Sroisuwan and contributors
# For license information, please see license.txt

"""Export every DocType to Excel, one workbook per module.

Each module workbook holds an ``_Index`` sheet (DocType / Category / Sheet /
Records) followed by one sheet per DocType. Data is fetched from a remote Frappe
site via HTTP using ``FrappeClient``, or from the local database when no client
is given.

DocType categorization is driven by a catalog loaded from
``DocType_Categorized.xlsx`` (the ``Catalog`` sheet produced by the
``cd_export_master`` project). When no catalog is provided, the system falls back
to the built-in classification based on ``issingle`` and ``is_submittable``.
"""

from __future__ import annotations

import io
import json
import re
from io import BytesIO
from pathlib import Path
from typing import Any, cast
from urllib.parse import urljoin

import frappe
import requests
from frappe.utils.xlsxutils import get_sanitized_sheet_name
from openpyxl import load_workbook

INDEX_SHEET_NAME = "_Index"
INDEX_HEADERS = ["DocType", "Category", "Sheet", "Records"]
MAX_SHEET_NAME_LENGTH = 31
EXCLUDED_CATALOG_CATEGORIES: frozenset[str] = frozenset({"Child Table", "Virtual DocType", "Review Required"})

ILLEGAL_XML_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
EXCEL_CELL_LIMIT = 32_767


class ExportError(RuntimeError):
    """Raised when an HTTP export cannot be completed."""


class FrappeClient:
    """Read-only REST client for a remote Frappe site, authenticated via API key/secret."""

    def __init__(
        self,
        base_url: str,
        api_key: str,
        api_secret: str,
        *,
        page_length: int = 200,
        timeout: float = 60,
        session: requests.Session | None = None,
    ) -> None:
        if not api_key or not api_secret:
            raise ExportError("FRAPPE_API_KEY and FRAPPE_API_SECRET are required")
        self.base_url = base_url.rstrip("/")
        self.page_length = page_length
        self.timeout = timeout
        self._session = session or requests.Session()
        self._session.headers.update(
            {
                "Authorization": f"token {api_key}:{api_secret}",
                "Accept": "application/json",
            }
        )

    def _json_get(self, path: str, *, params: dict[str, Any] | None = None) -> Any:
        url = f"{self.base_url}{path}"
        try:
            resp = self._session.get(url, params=params, timeout=self.timeout)
            resp.raise_for_status()
            return resp.json()
        except requests.RequestException as exc:
            _r = getattr(exc, "response", None)
            detail = f": {_r.text[:500]}" if _r is not None else ""
            raise ExportError(f"GET {url} failed{detail}") from exc
        except ValueError as exc:
            raise ExportError(f"GET {url} returned invalid JSON") from exc

    def _post(self, path: str, *, json_payload: dict[str, Any]) -> requests.Response:
        url = f"{self.base_url}{path}"
        try:
            resp = self._session.post(url, json=json_payload, timeout=self.timeout)
            resp.raise_for_status()
            return resp
        except requests.RequestException as exc:
            _r = getattr(exc, "response", None)
            detail = f": {_r.text[:500]}" if _r is not None else ""
            raise ExportError(f"POST {url} failed{detail}") from exc

    def _download_file(self, file_url: str) -> bytes:
        url = file_url if file_url.startswith(("http://", "https://")) else urljoin(self.base_url, file_url)
        try:
            resp = self._session.get(url, timeout=self.timeout)
            resp.raise_for_status()
            return resp.content
        except requests.RequestException as exc:
            _r = getattr(exc, "response", None)
            detail = f": {_r.text[:500]}" if _r is not None else ""
            raise ExportError(f"GET {url} failed{detail}") from exc

    def get_list(
        self,
        doctype: str,
        *,
        fields: list[str],
        filters: list[list[Any]] | None = None,
        order_by: str = "name asc",
    ) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        start = 0
        while True:
            payload = self._json_get(
                "/api/method/frappe.client.get_list",
                params={
                    "doctype": doctype,
                    "fields": json.dumps(fields),
                    "filters": json.dumps(filters or []),
                    "order_by": order_by,
                    "limit_start": start,
                    "limit_page_length": self.page_length,
                },
            )
            page = payload.get("message", [])
            if not isinstance(page, list):
                raise ExportError(f"Unexpected list response for {doctype}")
            rows.extend(page)
            if len(page) < self.page_length:
                return rows
            start += len(page)

    def export_data_workbook(self, doctype: str) -> bytes:
        """Export one DocType via Frappe's built-in Data Export API and return xlsx bytes."""
        import base64

        resp = self._post(
            "/api/method/frappe.core.doctype.data_export.exporter.export_data",
            json_payload={"doctype": doctype, "export_records": "all", "file_type": "Excel", "with_data": 1},
        )
        content_type = resp.headers.get("Content-Type", "").lower()
        if "spreadsheet" in content_type or resp.content[:2] == b"PK":
            return resp.content

        try:
            payload = resp.json()
        except ValueError as exc:
            raise ExportError(f"{doctype}: export_data returned non-Excel response") from exc

        message = payload.get("message")
        if isinstance(message, str):
            if message.startswith("/") or message.startswith(("http://", "https://")):
                return self._download_file(message)
            try:
                return base64.b64decode(message)
            except Exception:
                return message.encode("latin1")

        if isinstance(message, dict):
            for key in ("file_url", "fileurl"):
                file_url = message.get(key)
                if isinstance(file_url, str) and file_url:
                    return self._download_file(file_url)
            for key in ("file_content", "filecontent", "content"):
                content = message.get(key)
                if isinstance(content, str) and content:
                    try:
                        return base64.b64decode(content)
                    except Exception:
                        return content.encode("latin1")
                if isinstance(content, (bytes, bytearray)):
                    return bytes(content)

        raise ExportError(f"{doctype}: export_data response did not contain an Excel file")

    def get_count(self, doctype: str) -> int:
        """Return the record count for a DocType, or 0 on failure."""
        try:
            payload = self._json_get(
                "/api/method/frappe.client.get_count",
                params={"doctype": doctype},
            )
            return int(payload.get("message", 0))
        except Exception:
            return 0


def load_catalog(path: Path) -> dict[str, str]:
    """Load a DocType-to-category mapping from a ``DocType_Categorized.xlsx`` catalog file.

    The file must contain a ``Catalog`` sheet with at minimum the columns
    ``DocType``, ``Suggested Category``, and ``Category Override``. The effective
    category per DocType is the override when non-empty, otherwise the suggestion.
    """
    if not path.exists():
        raise ExportError(f"Catalog not found: {path}")
    wb = load_workbook(path, read_only=True, data_only=False)
    if "Catalog" not in wb.sheetnames:
        raise ExportError(f"{path} does not contain a 'Catalog' sheet")
    ws = wb["Catalog"]
    headers = [str(cell.value or "") for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    try:
        doctype_idx = headers.index("DocType")
        suggested_idx = headers.index("Suggested Category")
        override_idx = headers.index("Category Override")
    except ValueError as exc:
        raise ExportError(f"{path} is missing required Catalog columns: {exc}") from exc

    catalog: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        doctype = str(row[doctype_idx] or "").strip()
        suggested = str(row[suggested_idx] or "").strip()
        override = str(row[override_idx] or "").strip()
        if doctype:
            catalog[doctype] = override or suggested
    return catalog


def classify_doctype(*, issingle: bool, is_submittable: bool) -> str:
    """Categorize a DocType when no catalog is available.

    Single DocTypes are ``Setup``; submittable DocTypes are ``Transaction``;
    everything else is ``Master``.
    """
    if issingle:
        return "Setup"
    if is_submittable:
        return "Transaction"
    return "Master"


def list_target_doctypes(
    modules: tuple[str, ...] = (),
    *,
    client: FrappeClient | None = None,
    catalog: dict[str, str] | None = None,
) -> dict[str, list[dict[str, str]]]:
    """Group exportable DocTypes by module.

    When ``catalog`` is provided the effective category and exclusion list come
    from it (skipping ``Child Table``, ``Virtual DocType``, and
    ``Review Required``). Without a catalog, the built-in ``classify_doctype``
    logic is used and child/virtual DocTypes are excluded via DB filters.

    Pass ``client`` to fetch the DocType list from a remote site via HTTP.
    """
    if client:
        raw_rows = client.get_list(
            "DocType",
            fields=["name", "module", "istable", "issingle", "is_submittable", "is_virtual"],
            order_by="module asc, name asc",
        )
    else:
        filters: dict[str, object] = {"istable": 0, "is_virtual": 0}
        if modules:
            filters["module"] = ["in", list(modules)]
        raw_rows = frappe.get_all(
            "DocType",
            filters=filters,
            fields=["name", "module", "issingle", "is_submittable"],
            order_by="module asc, name asc",
        )

    grouped: dict[str, list[dict[str, str]]] = {}
    for row in raw_rows:
        name = cast(str, row["name"])
        module = cast(str, row.get("module") or "Uncategorized")

        if modules and module not in modules:
            continue

        if catalog is not None:
            category = catalog.get(name, "")
            if not category or category in EXCLUDED_CATALOG_CATEGORIES:
                continue
        else:
            if row.get("istable") or row.get("is_virtual"):
                continue
            category = classify_doctype(
                issingle=bool(row.get("issingle")),
                is_submittable=bool(row.get("is_submittable")),
            )

        grouped.setdefault(module, []).append({"name": name, "category": category})

    return grouped


def export_doctype_rows(doctype: str, *, client: FrappeClient | None = None) -> list[list[Any]]:
    """Return the data rows for a DocType as a list of rows (list of lists).

    When ``client`` is provided the rows are fetched via the remote Frappe Data
    Export API and the xlsx is parsed back to rows. Without a client, Frappe's
    local ``DataExporter`` is used instead.
    """
    if client:
        xlsx_bytes = client.export_data_workbook(doctype)
        return _workbook_bytes_to_rows(xlsx_bytes, doctype)

    import csv

    from frappe.core.doctype.data_export.exporter import DataExporter

    exporter = DataExporter(
        doctype=doctype, all_doctypes=True, with_data=True, file_type="CSV", template=False
    )
    try:
        exporter.build_response()
    except Exception:
        pass
    return list(csv.reader(io.StringIO(exporter.writer.getvalue())))


def _workbook_bytes_to_rows(xlsx_bytes: bytes, doctype: str) -> list[list[Any]]:
    try:
        wb = load_workbook(BytesIO(xlsx_bytes), data_only=False)
    except Exception as exc:
        raise ExportError(f"{doctype}: cannot parse Excel returned by export_data") from exc
    ws = wb.active
    rows: list[list[Any]] = [list(row) for row in ws.iter_rows(values_only=True)]
    return rows if rows else [["No data returned from export_data"]]


def _safe_count(doctype: str, *, client: FrappeClient | None = None) -> int:
    """Return the record count for a DocType from the remote API or local DB."""
    if client:
        return client.get_count(doctype)
    try:
        return frappe.db.count(doctype)
    except Exception:
        return 0


def _unique_sheet_name(doctype: str, used: set[str]) -> str:
    """Return an Excel-safe, unique sheet name derived from ``doctype``."""
    base = get_sanitized_sheet_name(doctype) or "Sheet"
    candidate = base
    suffix = 1
    while candidate in used:
        suffix += 1
        tail = f" {suffix}"
        candidate = base[: MAX_SHEET_NAME_LENGTH - len(tail)] + tail
    used.add(candidate)
    return candidate


def _sanitize_excel_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    value = ILLEGAL_XML_RE.sub("", value)
    if value.startswith(("=", "+", "-", "@")):
        value = f"'{value}"
    if len(value) > EXCEL_CELL_LIMIT:
        value = value[:EXCEL_CELL_LIMIT]
    return value


def build_module_workbook(
    doctypes: list[dict[str, str]],
    *,
    client: FrappeClient | None = None,
) -> bytes:
    """Build one xlsx (as bytes) with an ``_Index`` sheet plus a sheet per DocType.

    A DocType that fails to export is logged, listed in the index with an empty
    Sheet column, and skipped so one failure never aborts the whole workbook.
    """
    import xlsxwriter
    from frappe.utils.xlsxutils import make_xlsx

    used_names: set[str] = {INDEX_SHEET_NAME}
    index_rows: list[list[object]] = [list(INDEX_HEADERS)]
    sheets: list[tuple[str, list[list[Any]]]] = []

    for entry in doctypes:
        doctype = entry["name"]
        try:
            rows = export_doctype_rows(doctype, client=client)
        except Exception:
            frappe.logger(__name__).warning(f"Data export skipped DocType {doctype!r}", exc_info=True)
            index_rows.append([doctype, entry["category"], None, 0])
            continue

        sheet_name = _unique_sheet_name(doctype, used_names)
        sheets.append((sheet_name, rows))
        index_rows.append([doctype, entry["category"], sheet_name, _safe_count(doctype, client=client)])

    output = BytesIO()
    workbook = xlsxwriter.Workbook(output, {"constant_memory": True})
    try:
        make_xlsx(index_rows, INDEX_SHEET_NAME, wb=workbook)
        for sheet_name, rows in sheets:
            safe_rows = [[_sanitize_excel_value(cell) for cell in row] for row in rows]
            make_xlsx(safe_rows, sheet_name, wb=workbook)
    finally:
        workbook.close()
    return output.getvalue()
