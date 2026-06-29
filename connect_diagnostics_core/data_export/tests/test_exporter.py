# Copyright (c) 2026, Siwat Sroisuwan and contributors
# For license information, please see license.txt

from __future__ import annotations

import base64
import sys
import tempfile
from io import BytesIO
from pathlib import Path
from types import ModuleType
from typing import Any
from unittest import TestCase
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from connect_diagnostics_core.data_export import exporter
from connect_diagnostics_core.data_export.exporter import (
    EXCLUDED_CATALOG_CATEGORIES,
    INDEX_HEADERS,
    INDEX_SHEET_NAME,
    ExportError,
    FrappeClient,
    _sanitize_excel_value,
    _unique_sheet_name,
    _workbook_bytes_to_rows,
    build_module_workbook,
    classify_doctype,
    export_doctype_rows,
    list_target_doctypes,
    load_catalog,
)


def _xlsx_bytes(rows: list[list[Any]]) -> bytes:
    output = BytesIO()
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(output)
    return output.getvalue()


def _make_catalog_xlsx(rows: list[list[str]], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Catalog"
    for row in rows:
        ws.append(row)
    wb.save(path)


class TestClassifyDoctype(TestCase):
    def test_single_doctype_is_setup(self) -> None:
        self.assertEqual(classify_doctype(issingle=True, is_submittable=False), "Setup")

    def test_submittable_doctype_is_transaction(self) -> None:
        self.assertEqual(classify_doctype(issingle=False, is_submittable=True), "Transaction")

    def test_plain_doctype_is_master(self) -> None:
        self.assertEqual(classify_doctype(issingle=False, is_submittable=False), "Master")

    def test_single_takes_priority_over_submittable(self) -> None:
        self.assertEqual(classify_doctype(issingle=True, is_submittable=True), "Setup")


class TestUniqueSheetName(TestCase):
    def test_returns_doctype_name_when_unused(self) -> None:
        used: set[str] = set()
        self.assertEqual(_unique_sheet_name("Account", used), "Account")

    def test_deduplicates_collisions_with_suffix(self) -> None:
        used: set[str] = set()
        first = _unique_sheet_name("Account", used)
        second = _unique_sheet_name("Account", used)
        self.assertEqual(first, "Account")
        self.assertEqual(second, "Account 2")

    def test_truncates_to_excel_limit(self) -> None:
        used: set[str] = set()
        name = _unique_sheet_name("A" * 50, used)
        self.assertLessEqual(len(name), 31)

    def test_truncated_collision_stays_within_limit_and_unique(self) -> None:
        used: set[str] = set()
        long_name = "B" * 40
        first = _unique_sheet_name(long_name, used)
        second = _unique_sheet_name(long_name, used)
        self.assertLessEqual(len(second), 31)
        self.assertNotEqual(first, second)


class TestSanitizeExcelValue(TestCase):
    def test_formula_prefix_equals_is_escaped(self) -> None:
        self.assertEqual(_sanitize_excel_value("=SUM(A1)"), "'=SUM(A1)")

    def test_formula_prefix_plus_is_escaped(self) -> None:
        self.assertEqual(_sanitize_excel_value("+1"), "'+1")

    def test_formula_prefix_minus_is_escaped(self) -> None:
        self.assertEqual(_sanitize_excel_value("-1"), "'-1")

    def test_formula_prefix_at_is_escaped(self) -> None:
        self.assertEqual(_sanitize_excel_value("@cmd"), "'@cmd")

    def test_illegal_xml_chars_are_stripped(self) -> None:
        self.assertEqual(_sanitize_excel_value("hello\x00world\x08"), "helloworld")

    def test_string_truncated_at_cell_limit(self) -> None:
        long_str = "x" * 40_000
        result = _sanitize_excel_value(long_str)
        self.assertEqual(len(result), 32_767)

    def test_non_string_passes_through_unchanged(self) -> None:
        self.assertIsNone(_sanitize_excel_value(None))
        self.assertEqual(_sanitize_excel_value(42), 42)
        self.assertAlmostEqual(_sanitize_excel_value(3.14), 3.14)

    def test_normal_string_unchanged(self) -> None:
        self.assertEqual(_sanitize_excel_value("hello"), "hello")


class TestWorkbookBytesToRows(TestCase):
    def test_converts_xlsx_bytes_to_nested_list(self) -> None:
        expected = [["Header1", "Header2"], ["val1", "val2"]]
        rows = _workbook_bytes_to_rows(_xlsx_bytes(expected), "Test DocType")
        self.assertEqual(rows[0], ["Header1", "Header2"])
        self.assertEqual(rows[1], ["val1", "val2"])

    def test_invalid_bytes_raises_export_error(self) -> None:
        with self.assertRaises(ExportError):
            _workbook_bytes_to_rows(b"not an xlsx", "Bad DocType")


class TestLoadCatalog(TestCase):
    def test_override_wins_over_suggested_category(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "catalog.xlsx"
            _make_catalog_xlsx(
                [
                    ["DocType", "Suggested Category", "Category Override"],
                    ["Account", "Master Data", "Transaction"],
                ],
                path,
            )
            catalog = load_catalog(path)
        self.assertEqual(catalog["Account"], "Transaction")

    def test_suggested_used_when_override_is_empty(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "catalog.xlsx"
            _make_catalog_xlsx(
                [
                    ["DocType", "Suggested Category", "Category Override"],
                    ["Item", "Master Data", ""],
                ],
                path,
            )
            catalog = load_catalog(path)
        self.assertEqual(catalog["Item"], "Master Data")

    def test_blank_doctype_rows_are_skipped(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "catalog.xlsx"
            _make_catalog_xlsx(
                [
                    ["DocType", "Suggested Category", "Category Override"],
                    ["", "Master Data", ""],
                    ["Item", "Master Data", ""],
                ],
                path,
            )
            catalog = load_catalog(path)
        self.assertNotIn("", catalog)
        self.assertIn("Item", catalog)

    def test_missing_file_raises_export_error(self) -> None:
        with self.assertRaises(ExportError):
            load_catalog(Path("/nonexistent/catalog.xlsx"))

    def test_missing_catalog_sheet_raises_export_error(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "catalog.xlsx"
            wb = Workbook()
            wb.active.title = "NotCatalog"
            wb.save(path)
            with self.assertRaises(ExportError):
                load_catalog(path)

    def test_missing_required_column_raises_export_error(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "catalog.xlsx"
            _make_catalog_xlsx(
                [["DocType", "Suggested Category"]],  # missing Category Override
                path,
            )
            with self.assertRaises(ExportError):
                load_catalog(path)


class TestListTargetDoctypes(TestCase):
    @patch.object(exporter, "frappe")
    def test_groups_by_module_with_category(self, frappe: Mock) -> None:
        frappe.get_all.return_value = [
            {"name": "Account", "module": "Accounts", "issingle": 0, "is_submittable": 0},
            {"name": "Sales Invoice", "module": "Accounts", "issingle": 0, "is_submittable": 1},
            {"name": "Stock Settings", "module": "Stock", "issingle": 1, "is_submittable": 0},
        ]

        grouped = list_target_doctypes()

        self.assertEqual(set(grouped), {"Accounts", "Stock"})
        self.assertEqual(
            grouped["Accounts"],
            [
                {"name": "Account", "category": "Master"},
                {"name": "Sales Invoice", "category": "Transaction"},
            ],
        )
        self.assertEqual(grouped["Stock"], [{"name": "Stock Settings", "category": "Setup"}])

    @patch.object(exporter, "frappe")
    def test_excludes_child_and_virtual_doctypes(self, frappe: Mock) -> None:
        frappe.get_all.return_value = []

        list_target_doctypes()

        _, kwargs = frappe.get_all.call_args
        self.assertEqual(kwargs["filters"], {"istable": 0, "is_virtual": 0})

    @patch.object(exporter, "frappe")
    def test_module_filter_is_applied(self, frappe: Mock) -> None:
        frappe.get_all.return_value = []

        list_target_doctypes(modules=("Stock",))

        _, kwargs = frappe.get_all.call_args
        self.assertEqual(kwargs["filters"]["module"], ["in", ["Stock"]])

    @patch.object(exporter, "frappe")
    def test_catalog_uses_catalog_category_and_excludes_excluded_categories(self, frappe: Mock) -> None:
        frappe.get_all.return_value = [
            {
                "name": "Account",
                "module": "Accounts",
                "issingle": 0,
                "is_submittable": 0,
                "istable": 0,
                "is_virtual": 0,
            },
            {
                "name": "Item Price",
                "module": "Stock",
                "issingle": 0,
                "is_submittable": 0,
                "istable": 0,
                "is_virtual": 0,
            },
            {
                "name": "Journal Entry Detail",
                "module": "Accounts",
                "issingle": 0,
                "is_submittable": 0,
                "istable": 0,
                "is_virtual": 0,
            },
        ]
        catalog = {
            "Account": "Master Data",
            "Item Price": "Settings",
            "Journal Entry Detail": "Child Table",
        }

        grouped = list_target_doctypes(catalog=catalog)

        self.assertIn("Accounts", grouped)
        names_in_accounts = [d["name"] for d in grouped["Accounts"]]
        self.assertIn("Account", names_in_accounts)
        self.assertNotIn("Journal Entry Detail", names_in_accounts)
        self.assertEqual(grouped["Accounts"][0]["category"], "Master Data")

    @patch.object(exporter, "frappe")
    def test_catalog_excludes_all_excluded_catalog_categories(self, frappe: Mock) -> None:
        frappe.get_all.return_value = [
            {
                "name": f"DT {cat}",
                "module": "Core",
                "issingle": 0,
                "is_submittable": 0,
                "istable": 0,
                "is_virtual": 0,
            }
            for cat in EXCLUDED_CATALOG_CATEGORIES
        ]
        catalog = {f"DT {cat}": cat for cat in EXCLUDED_CATALOG_CATEGORIES}

        grouped = list_target_doctypes(catalog=catalog)

        self.assertEqual(grouped, {})

    def test_catalog_with_remote_client_uses_get_list(self) -> None:
        client = Mock(spec=FrappeClient)
        client.get_list.return_value = [
            {
                "name": "Payment Entry",
                "module": "Accounts",
                "issingle": 0,
                "is_submittable": 1,
                "istable": 0,
                "is_virtual": 0,
            },
        ]
        catalog = {"Payment Entry": "Transaction"}

        grouped = list_target_doctypes(client=client, catalog=catalog)

        client.get_list.assert_called_once()
        self.assertEqual(grouped["Accounts"][0]["category"], "Transaction")


class TestFrappeClientExportDataWorkbook(TestCase):
    def _make_client(self, session: Mock) -> FrappeClient:
        with patch.object(exporter.requests, "Session", return_value=session):
            return FrappeClient("https://example.test", "key", "secret")

    def test_returns_content_directly_when_magic_bytes_match(self) -> None:
        xlsx_content = _xlsx_bytes([["col"]])
        self.assertTrue(xlsx_content[:2] == b"PK")
        resp = Mock(headers={"Content-Type": "text/html"}, content=xlsx_content)
        resp.raise_for_status.return_value = None
        session = Mock()
        session.headers = {}
        session.post.return_value = resp

        client = self._make_client(session)
        result = client.export_data_workbook("Account")

        self.assertEqual(result, xlsx_content)

    def test_returns_content_when_spreadsheet_content_type(self) -> None:
        raw = b"PK\x03\x04fake_xlsx"
        resp = Mock(
            headers={"Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
            content=raw,
        )
        resp.raise_for_status.return_value = None
        session = Mock()
        session.headers = {}
        session.post.return_value = resp

        client = self._make_client(session)
        result = client.export_data_workbook("Account")

        self.assertEqual(result, raw)

    def test_follows_file_url_in_json_message(self) -> None:
        xlsx_content = _xlsx_bytes([["A"]])
        json_resp = Mock(headers={"Content-Type": "application/json"}, content=b"{}")
        json_resp.raise_for_status.return_value = None
        json_resp.content = b"\x00\x00"
        json_resp.json.return_value = {"message": "/files/export.xlsx"}

        file_resp = Mock(content=xlsx_content)
        file_resp.raise_for_status.return_value = None

        session = Mock()
        session.headers = {}
        session.post.return_value = json_resp
        session.get.return_value = file_resp

        client = self._make_client(session)
        result = client.export_data_workbook("Account")

        self.assertEqual(result, xlsx_content)

    def test_decodes_base64_in_json_message(self) -> None:
        xlsx_content = _xlsx_bytes([["A"]])
        encoded = base64.b64encode(xlsx_content).decode()
        resp = Mock(headers={"Content-Type": "application/json"}, content=b"\x00\x00")
        resp.raise_for_status.return_value = None
        resp.json.return_value = {"message": encoded}

        session = Mock()
        session.headers = {}
        session.post.return_value = resp

        client = self._make_client(session)
        result = client.export_data_workbook("Account")

        self.assertEqual(result, xlsx_content)

    def test_raises_export_error_when_no_xlsx_in_response(self) -> None:
        resp = Mock(headers={"Content-Type": "application/json"}, content=b"\x00\x00")
        resp.raise_for_status.return_value = None
        resp.json.return_value = {"message": {"no_file_key": "value"}}

        session = Mock()
        session.headers = {}
        session.post.return_value = resp

        client = self._make_client(session)
        with self.assertRaises(ExportError):
            client.export_data_workbook("Account")


class TestExportDoctypeRows(TestCase):
    def test_local_export_keeps_structure_when_frappe_stops_empty_export(self) -> None:
        module_name = "frappe.core.doctype.data_export.exporter"
        fake_module = ModuleType(module_name)

        class FakeDataExporter:
            def __init__(self, **kwargs) -> None:
                self.writer = Mock()
                self.writer.getvalue.return_value = (
                    "\nDocType:,Empty DocType\nColumn Labels:,ID,Title\nColumn Name:,name,title\n"
                )

            def build_response(self) -> None:
                self.data = []
                raise RuntimeError("No Data")

        fake_module.DataExporter = FakeDataExporter  # ty: ignore[unresolved-attribute] - dynamically injected into fake ModuleType for test
        original_module = sys.modules.get(module_name)
        sys.modules[module_name] = fake_module
        try:
            rows = export_doctype_rows("Empty DocType")
        finally:
            if original_module is None:
                del sys.modules[module_name]
            else:
                sys.modules[module_name] = original_module

        self.assertEqual(rows[1], ["DocType:", "Empty DocType"])
        self.assertEqual(rows[3], ["Column Name:", "name", "title"])

    def test_remote_export_calls_export_data_workbook_and_parses_rows(self) -> None:
        xlsx_content = _xlsx_bytes([["DocType:", "Account"], ["name", "account_name"]])
        client = Mock(spec=FrappeClient)
        client.export_data_workbook.return_value = xlsx_content

        rows = export_doctype_rows("Account", client=client)

        client.export_data_workbook.assert_called_once_with("Account")
        self.assertEqual(rows[0], ["DocType:", "Account"])
        self.assertEqual(rows[1], ["name", "account_name"])


class TestBuildModuleWorkbook(TestCase):
    @patch.object(exporter, "frappe")
    @patch.object(exporter, "export_doctype_rows")
    def test_workbook_has_index_and_doctype_sheets(self, export_doctype_rows: Mock, frappe: Mock) -> None:
        frappe.db.count.return_value = 3
        export_doctype_rows.return_value = [
            [""],
            ["DocType:", "Account"],
            ["Column Labels:", "ID", "Account Name"],
            ["Column Name:", "name", "account_name"],
            ["Mandatory:", "Yes", "Yes"],
            ["Type:", "Data", "Data"],
            ["Info:", "", ""],
        ]

        content = build_module_workbook([{"name": "Account", "category": "Master"}])

        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook.sheetnames, [INDEX_SHEET_NAME, "Account"])

        index = workbook[INDEX_SHEET_NAME]
        self.assertEqual([cell.value for cell in index[1]], INDEX_HEADERS)
        self.assertEqual([cell.value for cell in index[2]], ["Account", "Master", "Account", 3])

        account = workbook["Account"]
        self.assertEqual(account["A2"].value, "DocType:")
        self.assertEqual(account["B2"].value, "Account")
        self.assertEqual(account["A4"].value, "Column Name:")

    @patch.object(exporter, "frappe")
    @patch.object(exporter, "export_doctype_rows")
    def test_empty_doctype_gets_sheet_and_structure(self, export_doctype_rows: Mock, frappe: Mock) -> None:
        frappe.db.count.return_value = 0
        export_doctype_rows.return_value = [
            [""],
            ["DocType:", "Empty DocType"],
            ["Column Labels:", "ID", "Title"],
            ["Column Name:", "name", "title"],
            ["Mandatory:", "Yes", "No"],
            ["Type:", "Data", "Data"],
            ["Info:", "", ""],
        ]

        content = build_module_workbook([{"name": "Empty DocType", "category": "Master"}])

        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook.sheetnames, [INDEX_SHEET_NAME, "Empty DocType"])

        index = workbook[INDEX_SHEET_NAME]
        self.assertEqual([cell.value for cell in index[2]], ["Empty DocType", "Master", "Empty DocType", 0])

        worksheet = workbook["Empty DocType"]
        self.assertEqual(worksheet["A4"].value, "Column Name:")
        self.assertEqual(worksheet["B4"].value, "name")

    @patch.object(exporter, "frappe")
    @patch.object(exporter, "export_doctype_rows")
    def test_failed_doctype_is_indexed_but_skipped(self, export_doctype_rows: Mock, frappe: Mock) -> None:
        export_doctype_rows.side_effect = Exception("no table")

        content = build_module_workbook([{"name": "System Settings", "category": "Setup"}])

        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook.sheetnames, [INDEX_SHEET_NAME])

        index = workbook[INDEX_SHEET_NAME]
        self.assertEqual([cell.value for cell in index[2]], ["System Settings", "Setup", None, 0])
